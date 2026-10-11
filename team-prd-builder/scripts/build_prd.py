import argparse
import hashlib
import json
import re
import shutil
import tempfile
from pathlib import Path

from diagram_export import diagram_png, diagram_svg, diagram_xml, label_size
from validate_prd import DEFAULT_KINDS, validate
from spec_tests import run_suite


STATUS_LABELS = {
    "confirmed": "已確認", "proposed": "建議待審", "pending": "待確認", "excluded": "不在本期",
    "superseded": "已取代", "open": "未決", "answered": "已回答", "approved": "已核准",
    "not_run": "未執行", "passed": "有通過紀錄", "failed": "有失敗紀錄",
    "normal": "正常", "rejection": "拒絕與異常", "permission": "權限",
    "observed": "僅觀察", "author_confirmed": "作者已驗證", "read": "已讀取", "unavailable": "尚未取得",
}


def heading_text(value):
    return re.sub(r"[^\w\s]", " ", str(value)).strip()


def document(model, destination, report, docx_font="Arial Unicode MS"):
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION_START
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.title = model["project"]["title"]
    doc.core_properties.subject = "團隊開發產品需求與圖解"
    for style_name, size in (("Normal", 11), ("Title", 24), ("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12), ("Caption", 10)):
        style = doc.styles[style_name]
        style.font.name = docx_font
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attribute), docx_font)
            fonts.attrib.pop(qn("w:" + attribute + "Theme"), None)
        language = OxmlElement("w:lang")
        language.set(qn("w:val"), "zh-TW")
        language.set(qn("w:eastAsia"), "zh-TW")
        style.element.get_or_add_rPr().append(language)
        style.paragraph_format.space_after = Pt(4)
        if style_name == "Normal":
            style.paragraph_format.line_spacing = 1.15
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[style_name].paragraph_format.space_before = Pt(10)
        doc.styles[style_name].paragraph_format.keep_with_next = True

    def orientation(section, landscape=False):
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width = Cm(29.7 if landscape else 21)
        section.page_height = Cm(21 if landscape else 29.7)
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.footer_distance = Cm(0.8)

    orientation(doc.sections[0])
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    page_field.append(field_run)
    footer.paragraph_format.space_after = Pt(0)
    bookmark_counter = 0

    def paragraph(text="", label=None, keep=False):
        item = doc.add_paragraph()
        item.paragraph_format.keep_with_next = keep
        if label:
            item.add_run(label + " ").bold = True
        item.add_run(str(text))
        return item

    def heading(text, level=1, identifier=None):
        nonlocal bookmark_counter
        item = doc.add_heading(heading_text(text), level=level)
        if identifier:
            bookmark_counter += 1
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(bookmark_counter))
            start.set(qn("w:name"), identifier.replace("-", "_"))
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(bookmark_counter))
            item._p.insert(0, start)
            item._p.append(end)
        return item

    def bullets(values):
        for value in values:
            paragraph(value).paragraph_format.left_indent = Cm(0.3)

    def table(headers, rows, widths):
        element = doc.add_table(rows=1, cols=len(headers))
        element.autofit = False
        for column, width in zip(element.columns, widths):
            column.width = Cm(width)
        for cell, value, width in zip(element.rows[0].cells, headers, widths):
            cell.width = Cm(width)
            cell.text = value
        repeat = OxmlElement("w:tblHeader")
        element.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = element.add_row().cells
            for cell, value, width in zip(cells, row, widths):
                cell.width = Cm(width)
                cell.text = str(value)
        for row_index, row in enumerate(element.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for direction in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + direction)
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "4")
                    border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                properties.append(borders)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "17324F" if row_index == 0 else ("F3F6FA" if row_index % 2 else "FFFFFF"))
                properties.append(shading)
                margins = OxmlElement("w:tcMar")
                for direction in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + direction)
                    margin.set(qn("w:w"), "100")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                for item in cell.paragraphs:
                    item.paragraph_format.space_after = Pt(3)
                    for run in item.runs:
                        run.font.size = Pt(10)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        paragraph()

    def feature_card(feature):
        heading(feature["id"] + " " + feature["title"], 2, feature["id"])
        paragraph(feature["story"])
        paragraph(STATUS_LABELS[feature["status"]] + "  " + feature["priority"], "規格狀態")
        for key, label in (("roles", "角色"), ("preconditions", "前置條件"), ("inputs", "輸入"), ("steps", "操作步驟"), ("success", "成功結果"), ("errors", "拒絕與異常"), ("recovery", "復原")):
            value = feature[key]
            if isinstance(value, list):
                if value:
                    paragraph("；".join(value), label)
            else:
                paragraph(value, label)
        paragraph("；".join(feature["screen_ids"]) or "未指定", "畫面")
        paragraph("；".join(feature.get("mockup_ids", [])) or "未提供Mockup，不阻擋本功能撰寫", "Mockup")
        for key, label in (("rule_ids", "共用規則"), ("diagram_ids", "圖"), ("acceptance_ids", "驗收"), ("source_refs", "來源")):
            paragraph("；".join(feature[key]) or "無", label)
        if feature.get("decision_id"):
            paragraph(feature["decision_id"], "確認決策")
        blockers = [item["id"] for item in model["questions"] if item["status"] == "open" and feature["id"] in item["blocking_feature_ids"]]
        if blockers:
            paragraph("；".join(blockers), "未決阻擋")
        if feature.get("delivery"):
            for key, label in (("engineering_review", "工程審閱"), ("implementation", "實作"), ("testing", "產品測試"), ("deployment", "部署")):
                paragraph(feature["delivery"][key] + "  證據 " + feature["delivery"].get(key + "_evidence", "未提供"), label)

    def mockup_records():
        heading("Mockup 與需求對照", 2)
        if not model.get("mockups"):
            paragraph("本文件未使用HTML Mockup。需求仍依文字、訪談及來源確認撰寫，缺少原型不構成交付阻擋。")
        for item in model.get("mockups", []):
            heading(item["id"] + " " + item["title"], 3, item["id"])
            paragraph(item["version"] + "  " + item["locator"], "版本及定位")
            paragraph(STATUS_LABELS[item["availability"]] + "  " + STATUS_LABELS[item["confirmation"]], "判讀狀態")
            paragraph("；".join(item["feature_ids"]), "對應功能")
            paragraph("；".join(item["screen_ids"]), "畫面")
            bullets(item["observations"])
            paragraph("；".join(item["confirmed_scope"]) or "未記錄", "已驗證範圍")
            if item.get("confirmation_evidence"):
                paragraph(item["confirmation_evidence"], "確認證據")
            if item.get("confirmed_by"):
                paragraph(item["confirmed_by"] + "  " + item.get("confirmed_on", ""), "確認人及日期")
            paragraph("；".join(item["confirmed_decision_ids"]) or "無", "對應決策")
            bullets(item["limitations"])

    doc.add_paragraph(heading_text(model["project"]["title"]), style="Title")
    paragraph(model["project"]["summary"])
    paragraph(model["project"]["date"] + "  " + model["project"]["profile"], "文件資訊")
    paragraph("本文件供產品、工程與QA共同定義範圍、可觀察行為及驗收。請依來源與各項確認狀態判斷適用性；文件產出不代表功能已實作或產品測試通過。")
    paragraph(model["project"]["disclaimer"])
    heading("五層閱讀導覽", 1)
    for level in sorted(model["explanation_levels"], key=lambda item: item["level"]):
        paragraph("L" + str(level["level"]) + "  " + level["title"] + "  " + "、".join(level["chapter_ids"]), keep=True)
        paragraph(level["note"])
    doc.add_page_break()
    features = {item["id"]: item for item in model["features"]}
    rules = {item["id"]: item for item in model["rules"]}
    diagrams = {item["id"]: item for item in model["diagrams"]}
    readings = []
    for chapter in model["chapters"]:
        kind = chapter.get("kind", DEFAULT_KINDS.get(chapter["id"], "general"))
        chapter_diagrams = [diagrams[identifier] for identifier in chapter["diagram_ids"]]
        landscape = bool(chapter_diagrams and chapter_diagrams[0].get("orientation") == "landscape")
        if chapter_diagrams:
            orientation(doc.add_section(WD_SECTION_START.NEW_PAGE), landscape)
        elif doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE:
            orientation(doc.add_section(WD_SECTION_START.NEW_PAGE))
        elif kind == "features":
            doc.add_page_break()
        heading(chapter["id"] + " " + chapter["title"], 1, chapter["id"])
        if chapter["summary"]:
            paragraph(chapter["summary"], keep=bool(chapter_diagrams))
        for text in chapter["paragraphs"]:
            paragraph(text, keep=bool(chapter_diagrams))
        if kind == "summary":
            heading("需求來源", 2)
            for item in model["sources"]:
                paragraph(item["id"] + "  " + item["title"], keep=True)
                paragraph(item["kind"] + "  " + item["locator"] + "  權威來源 " + ("是" if item["authoritative"] else "否"))
        if kind == "scope":
            heading("本期範圍", 2)
            bullets(model["project"]["mvp"])
            heading("不做範圍", 2)
            bullets(model["project"]["out_of_scope"])
        for index, diagram in enumerate(chapter_diagrams):
            if index:
                orientation(doc.add_section(WD_SECTION_START.NEW_PAGE), diagram.get("orientation") == "landscape")
            paragraph(diagram["purpose"], "閱讀目的", keep=True)
            caption = doc.add_paragraph(diagram["id"] + "  " + diagram["title"], style="Caption")
            caption.paragraph_format.keep_with_next = True
            current_landscape = diagram.get("orientation") == "landscape"
            maximum_width, maximum_height = (25.7, 11) if current_landscape else (17, 17)
            if current_landscape and diagram["type"] == "state":
                maximum_height = 9.5
            picture_width = min(maximum_width, maximum_height * diagram["width"] / diagram["height"])
            smallest = min([diagram.get("font_size", 14)] + [label_size(diagram, edge) for edge in diagram["edges"] if edge.get("label")])
            effective_points = smallest * picture_width / diagram["width"] * 72 / 2.54
            if effective_points < 8:
                raise ValueError(diagram["id"] + " 縮入Word後約" + str(round(effective_points, 1)) + "pt；請增字級、改方向或拆圖。")
            image_paragraph = doc.add_paragraph()
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_paragraph.paragraph_format.keep_with_next = True
            image_run = image_paragraph.add_run()
            picture = image_run.add_picture(str(destination.parent / "diagrams" / (diagram["id"] + ".png")), width=Cm(picture_width))
            picture._inline.docPr.set("descr", diagram["purpose"])
            for text, label in (("；".join(diagram["reading_notes"]), "圖解"), ("；".join(diagram["legend"]), "圖例"), ("；".join(diagram["feature_ids"]) + "  未決 " + ("；".join(diagram["question_ids"]) or "無") + "  原檔 drawio/" + diagram["id"] + ".drawio", "追溯")):
                note = paragraph(text, label, keep=label != "追溯")
                note.paragraph_format.space_after = Pt(3)
                note.paragraph_format.line_spacing = 1
                for run in note.runs:
                    run.font.size = Pt(10)
            readings.append({"id": diagram["id"], "approx_minimum_label_points": round(effective_points, 1), "orientation": diagram.get("orientation", "portrait")})
        if chapter_diagrams and (chapter["rule_ids"] or chapter["feature_ids"]):
            orientation(doc.add_section(WD_SECTION_START.NEW_PAGE))
            heading(chapter["title"] + " 詳述", 2)
        for identifier in chapter["rule_ids"]:
            item = rules[identifier]
            heading(item["id"] + " " + item["title"], 2, item["id"])
            paragraph(STATUS_LABELS[item["status"]] + "  " + item["kind"] + "  生效階段 " + item["stage"], keep=True)
            paragraph(item["statement"])
            paragraph("；".join(item["feature_ids"]) + "  來源 " + "；".join(item["source_refs"]), "適用範圍")
        for feature_index, identifier in enumerate(chapter["feature_ids"]):
            if feature_index:
                doc.add_page_break()
            feature_card(features[identifier])
        if kind == "features":
            mockup_records()
        if kind == "quality":
            heading("Given When Then 驗收案例", 2)
            for item in model["acceptance"]:
                heading(item["id"] + " " + item["feature_id"] + " " + STATUS_LABELS.get(item["category"], item["category"]), 3, item["id"])
                for key, label in (("given", "Given"), ("when", "When"), ("then", "Then")):
                    paragraph(item[key], label, keep=True)
                paragraph(STATUS_LABELS[item["execution"]] + "  證據 " + (item.get("evidence") or "未提供"), "執行")
            heading("跨職能審閱", 2)
            table(["角色", "狀態", "證據"], [(item["role"], STATUS_LABELS[item["status"]], item.get("evidence") or "未提供") for item in model["reviews"]], [3, 3, 11])
        if kind == "governance":
            heading("變更提案與確認", 2)
            if not model["changes"]:
                paragraph("本文件未記錄業務變更。這不代表已完成等義審查。重構時必須保留差異及PM決策。")
            for change in model["changes"]:
                heading(change["id"] + " " + change["kind"], 3)
                paragraph(change["status"] + "  " + "、".join(change["scope_ids"]), "狀態與影響")
                paragraph(change["before"], "原文")
                paragraph(change["after"], "提案")
                paragraph(change["reason"], "理由")
                paragraph(change.get("decision_id", "待PM確認，未套用"), "決策")
            heading("待決問題", 2)
            for item in model["questions"]:
                heading(item["id"] + " " + item["title"], 3, item["id"])
                paragraph(item["context"])
                paragraph(item["owner"] + "  " + STATUS_LABELS[item["status"]] + "  阻擋 " + ("；".join(item["blocking_feature_ids"]) or "無"), "責任與影響")
                paragraph("；".join(item["options"]) or "無", "選項")
                paragraph(item.get("answer") or "未回答", "結論")
            heading("確認決策", 2)
            for item in model["decisions"]:
                paragraph(item["id"] + "  " + item["title"], keep=True)
                paragraph(item["answer"])
                paragraph(item["owner"] + "  " + item["date"] + "  範圍 " + "；".join(item["scope_ids"]) + "  來源 " + "；".join(item["source_refs"]))
            if not model["decisions"]:
                paragraph("尚未記錄確認決策。")
            heading("原文映射", 2)
            if model["mappings"]:
                table(["來源及原章節", "對應功能"], [(item["source_ref"] + " / " + item["original_id"], "；".join(item["feature_ids"])) for item in model["mappings"]], [6, 11])
            else:
                paragraph("未使用舊章節重構映射。")
    orientation(doc.add_section(WD_SECTION_START.NEW_PAGE))
    heading("目標與機器驗證", 1)
    paragraph("以下測試只執行規格參考模型。它不是產品實作。真實系統測試尚未執行。PM尚未確認的預期不得列為正式通過。")
    table(["目標", "功能", "測試及AC"], [(goal["id"] + " " + goal["title"], "、".join(goal["feature_ids"]), "\n".join(case["id"] + "／" + "、".join(case["acceptance_ids"]) for case in model["checks"] if case["goal_id"] == goal["id"]) or "缺口，待PM確認") for goal in model["goals"]], [5, 3, 9])
    for key, title in (("spec_validation", "正式規格測試"), ("example_validation", "合成案例展示")):
        if key not in report:
            continue
        result = report[key]
        heading(title, 2)
        paragraph(result["status"] + "  passed " + str(result["counts"]["passed"]) + "  failed " + str(result["counts"]["failed"]) + "  blocked " + str(result["counts"]["blocked_pending_pm"]), "結果")
        command = paragraph("python validation/run_spec_tests.py" + (" --example --output validation/example-results.json" if key == "example_validation" else " --output validation/results.json"), "命令")
        for run in command.runs:
            run.font.size = Pt(9)
        for label, filename in (("契約", "contracts.json"), ("獨立預期", "cases.json"), ("執行器", "run_spec_tests.py")):
            paragraph("validation/" + filename, label)
        if key == "example_validation":
            paragraph("example_only代表合成案例通過。它不核准任何需求，也不計入正式通過。")
        for entry in result["cases"]:
            if key == "example_validation":
                paragraph(entry["id"] + "  " + entry["status"] + "  expected " + json.dumps(entry["expected"], ensure_ascii=False) + "  actual " + json.dumps(entry.get("actual"), ensure_ascii=False))
    heading("未驗證範圍", 2)
    bullets(model["project"].get("validation_gaps", ["尚未記錄驗證限制，待PM確認。不能聲稱全部功能已驗證。"] ))
    for item in list(doc.paragraphs) + [entry for element in doc.tables for row in element.rows for cell in row.cells for entry in cell.paragraphs]:
        for run in item.runs:
            fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + attribute), docx_font)
                fonts.attrib.pop(qn("w:" + attribute + "Theme"), None)
    doc.save(destination)
    report["docx_generation"] = {"status": "generated", "embedded_figures": len(readings), "readability_estimates": readings, "note": "字級估算不是逐頁渲染與人工目視。"}
    return doc


def build(model, output, handoff=False, overwrite=False, font_file=None, docx_font="Arial Unicode MS", example_tests=False):
    if not isinstance(docx_font, str) or not docx_font.strip():
        raise ValueError("DOCX字型名稱不能留空。")
    report = validate(model, handoff)
    if report["errors"]:
        raise ValueError(json.dumps(report, ensure_ascii=False, indent=2))
    output = Path(output)
    if output.exists() and any(output.iterdir()) and not overwrite:
        raise FileExistsError("輸出已有內容；確認重生成才加 --overwrite：" + str(output))
    if output.exists() and ((output / "index.html").exists() or (output / "prd.md").exists()):
        raise ValueError("此目錄含舊版HTML交付，請使用新目錄，避免混入已淘汰輸出。")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".prd-build-", dir=output.parent) as directory:
        staging = Path(directory)
        (staging / "diagrams").mkdir()
        (staging / "drawio").mkdir()
        validation_path = staging / "validation"
        validation_path.mkdir()
        context = {key: model[key] for key in ("sources", "decisions", "goals", "features", "rules", "contracts", "questions", "changes")}
        for name, contents in (("contracts.json", context), ("cases.json", model["checks"])):
            (validation_path / name).write_text(json.dumps(contents, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
        shutil.copy2(Path(__file__).with_name("spec_tests.py"), validation_path / "run_spec_tests.py")
        for example, key, filename in [(False, "spec_validation", "results.json")] + ([(True, "example_validation", "example-results.json")] if example_tests else []):
            result = run_suite(context, model["checks"], example)
            result["input_sha256"] = {name: hashlib.sha256((validation_path / name).read_bytes()).hexdigest() for name in ("contracts.json", "cases.json")}
            report[key] = result
            (validation_path / filename).write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
        if handoff and report["spec_validation"]["status"] != "passed":
            raise ValueError("正式規格測試未通過；不能列為開發交付。")
        for diagram in model["diagrams"]:
            (staging / "diagrams" / (diagram["id"] + ".svg")).write_text(diagram_svg(diagram), encoding="utf-8")
            diagram_png(diagram, staging / "diagrams" / (diagram["id"] + ".png"), font_file)
            (staging / "drawio" / (diagram["id"] + ".drawio")).write_text(diagram_xml([diagram]), encoding="utf-8")
        (staging / "diagrams.drawio").write_text(diagram_xml(model["diagrams"]), encoding="utf-8")
        (staging / "prd.json").write_text(json.dumps(model, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
        document(model, staging / "prd.docx", report, docx_font)
        (staging / "quality-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        for file in staging.rglob("*"):
            if file.is_file():
                destination = output / file.relative_to(staging)
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(file, destination)
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--handoff", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--font-file", type=Path)
    parser.add_argument("--docx-font", default="Arial Unicode MS")
    parser.add_argument("--example-tests", action="store_true")
    args = parser.parse_args()
    try:
        model = json.loads(args.input.read_text(encoding="utf-8"))
        report = build(model, args.output, args.handoff, args.overwrite, args.font_file, args.docx_font, args.example_tests)
    except (OSError, ValueError, KeyError, TypeError, AttributeError, ImportError) as exception:
        print(str(exception))
        return 1
    print(json.dumps({"output": str(args.output.resolve()), "structure": report["structure"], "handoff_candidates": report["handoff_candidates"], "embedded_figures": report["docx_generation"]["embedded_figures"], "render_review": "not_run"}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
